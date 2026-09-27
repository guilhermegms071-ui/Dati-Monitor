"""Park screen (PROMPT 10.6), device detail, dashboard and the disconnected-devices job."""

import io
import uuid
from datetime import UTC, date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

import httpx
from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.models import Agent, AuditLog, Customer, Device, DeviceEvent, Reading, ReadingAdjustment, Site
from app.services import park as park_svc
from tests.agent_helpers import FakeAgent, enrolled_agent
from tests.conftest import Factory, Tenant, auth, login


def supply(
    color: str, percent: float | None, key: str, *, state: str = "ok", kind: str = "toner"
) -> dict[str, Any]:
    return {
        "key": key,
        "description": f"Toner ({color})",
        "type": kind,
        "class": "consumed",
        "color": color,
        "level": None if percent is None else int(percent),
        "max_capacity": 100,
        "percent": percent,
        "level_state": state,
        "unit": "percent",
    }


SP = ZoneInfo("America/Sao_Paulo")
CRLF = chr(13) + chr(10)
LAST_READ = datetime.now(UTC) - timedelta(minutes=1)


def sp_noon(days_before: int) -> datetime:
    """Noon (São Paulo) N days before the day of LAST_READ: always a different day in the series."""
    day: date = LAST_READ.astimezone(SP).date() - timedelta(days=days_before)
    return datetime.combine(day, time(12, 0), tzinfo=SP).astimezone(UTC)


async def populate(agent: FakeAgent) -> None:
    """Three printers of the same agent: two colors with supplies, one mono; different counters."""
    now = LAST_READ
    items = [
        agent.reading(
            "KM-001", {"total": 217031, "mono": 100150, "color": 116881}, ip="10.0.0.11", read_at=now
        ),
        agent.item(
            "supplies",
            serial="KM-001",
            ip="10.0.0.11",
            supplies=[
                supply("black", 25, "1.1"),
                supply("cyan", 55, "1.2"),
                supply("magenta", 66, "1.3"),
                supply("yellow", 5, "1.4"),
            ],
        ),
        agent.reading(
            "CAN-002", {"total": 150000, "mono": 90000, "color": 60000}, ip="10.0.0.12", read_at=now
        ),
        agent.item(
            "supplies",
            serial="CAN-002",
            ip="10.0.0.12",
            supplies=[supply("black", None, "1.1", state="unknown")],
        ),
        agent.reading("MONO-003", {"total": 45678, "mono": 45678}, ip="10.0.0.13", read_at=now),
    ]
    results = await agent.send(items)
    assert all(r["status"] == "accepted" for r in results), results


async def setup(client: httpx.AsyncClient, factory: Factory) -> tuple[Tenant, str, FakeAgent]:
    tenant = await factory.tenant()
    admin = await login(client, tenant.admin_email)
    agent = await enrolled_agent(client, tenant)
    await agent.heartbeat()  # vira MASTER do local
    await populate(agent)
    return tenant, admin, agent


async def park(client: httpx.AsyncClient, token: str, **params: Any) -> dict[str, Any]:
    resp = await client.get("/api/v1/park", params=params, headers=auth(token))
    assert resp.status_code == 200, resp.text
    data: dict[str, Any] = resp.json()
    return data


async def test_park_rows_filters_sort_and_levels(client: httpx.AsyncClient, factory: Factory) -> None:
    tenant, admin, _ = await setup(client, factory)
    data = await park(client, admin, sort="total", direction="desc")
    assert data["total"] == 3
    assert [r["serial"] for r in data["items"]] == ["KM-001", "CAN-002", "MONO-003"]
    km = data["items"][0]
    assert km["customer_name"] == "Revenda A Cliente"
    assert km["site_name"] == "Revenda A Local"
    assert km["agent_name"] == "Coletor 1"
    assert km["last_mono"] == 100150
    assert km["last_color"] == 116881
    # Barras na ordem C, M, Y, K.
    assert [(s["color"], s["percent"]) for s in km["supplies"]] == [
        ("cyan", 55.0),
        ("magenta", 66.0),
        ("yellow", 5.0),
        ("black", 25.0),
    ]
    can = data["items"][1]
    assert can["supplies"] == [
        {"color": "black", "percent": None, "level_state": "unknown", "description": "Toner (black)"}
    ]

    assert [r["serial"] for r in (await park(client, admin, q="can-0"))["items"]] == ["CAN-002"]
    assert (await park(client, admin, q="revenda a cliente"))["total"] == 3
    assert [r["serial"] for r in (await park(client, admin, ip="10.0.0.13"))["items"]] == ["MONO-003"]
    assert (await park(client, admin, status=["error"]))["total"] == 0
    assert (await park(client, admin, status=["ready", "error"]))["total"] == 3
    assert (await park(client, admin, customer_id=str(tenant.customer_id)))["total"] == 3
    assert (await park(client, admin, customer_id=str(uuid.uuid4())))["total"] == 0

    # Paginação por cursor (a tabela virtualizada carrega em blocos).
    first = await park(client, admin, sort="serial", limit=2)
    assert len(first["items"]) == 2
    assert first["next_cursor"]
    second = await park(client, admin, sort="serial", limit=2, cursor=first["next_cursor"])
    assert [r["serial"] for r in first["items"] + second["items"]] == ["CAN-002", "KM-001", "MONO-003"]
    assert second["next_cursor"] is None
    for sort in park_svc.PARK_SORTS:
        assert (
            await client.get("/api/v1/park", params={"sort": sort}, headers=auth(admin))
        ).status_code == 200
    assert (
        await client.get("/api/v1/park", params={"sort": "senha"}, headers=auth(admin))
    ).status_code == 400


async def test_park_is_scoped(client: httpx.AsyncClient, factory: Factory) -> None:
    tenant, _, _ = await setup(client, factory)
    other = await factory.tenant("Revenda B")
    other_admin = await login(client, other.admin_email)
    assert (await park(client, other_admin))["total"] == 0
    # Usuário restrito a outro cliente da mesma revenda não vê nada.
    maker = factory.maker
    async with maker() as s:
        c2 = Customer(reseller_id=tenant.reseller_id, company_id=tenant.company_id, name="Outro cliente")
        s.add(c2)
        await s.commit()
    _, email = await factory.user(tenant.reseller_id, role="customer_viewer", customer_id=c2.id)
    viewer = await login(client, email)
    assert (await park(client, viewer))["total"] == 0
    _, email = await factory.user(tenant.reseller_id, role="customer_viewer", customer_id=tenant.customer_id)
    viewer = await login(client, email)
    assert (await park(client, viewer))["total"] == 3


async def test_update_move_bulk_and_events(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    tenant, admin, agent = await setup(client, factory)
    rows = {r["serial"]: r for r in (await park(client, admin))["items"]}
    km = rows["KM-001"]["id"]
    resp = await client.patch(
        f"/api/v1/devices/{km}", json={"asset_tag": " PAT-77 ", "sector": "Financeiro"}, headers=auth(admin)
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["asset_tag"] == "PAT-77"
    assert resp.json()["sector"] == "Financeiro"

    # Mover para outro local do mesmo cliente.
    async with sessionmaker() as s:
        site2 = Site(reseller_id=tenant.reseller_id, customer_id=tenant.customer_id, name="Filial 2")
        s.add(site2)
        await s.commit()
    resp = await client.patch(f"/api/v1/devices/{km}", json={"site_id": str(site2.id)}, headers=auth(admin))
    assert resp.json()["site_name"] == "Filial 2"
    bad = await client.patch(
        f"/api/v1/devices/{km}", json={"site_id": str(uuid.uuid4())}, headers=auth(admin)
    )
    assert bad.status_code == 400

    ids = [r["id"] for r in rows.values()]
    resp = await client.post(
        "/api/v1/devices/bulk",
        json={"device_ids": ids, "action": "update", "sector": "TI"},
        headers=auth(admin),
    )
    assert resp.json()["changed"] == 3
    resp = await client.post(
        "/api/v1/devices/bulk",
        json={"device_ids": [km, str(uuid.uuid4())], "action": "deactivate"},
        headers=auth(admin),
    )
    body = resp.json()
    assert body["changed"] == 1
    assert body["skipped"][0]["reason"] == "não encontrado"
    assert (await park(client, admin))["total"] == 2
    assert (await park(client, admin, inactive=True))["total"] == 1
    counts = (await client.get("/api/v1/park/counts", headers=auth(admin))).json()
    assert counts == {"total": 2, "disconnected": 0, "inactive": 1}

    # "Ler agora" em massa vira um read_now para o MASTER do local, com os equipamentos escolhidos.
    active = [r["id"] for r in (await park(client, admin))["items"]]
    resp = await client.post(
        "/api/v1/devices/bulk", json={"device_ids": active, "action": "read_now"}, headers=auth(admin)
    )
    assert resp.status_code == 200, resp.text
    cmd_ids = resp.json()["commands"]
    assert len(cmd_ids) >= 1
    cmd = (await client.get(f"/api/v1/commands/{cmd_ids[0]}", headers=auth(admin))).json()
    assert cmd["type"] == "read_now"
    assert cmd["agent_id"] == agent.agent_id

    async with sessionmaker() as s:
        types = set(
            (
                await s.execute(select(DeviceEvent.type).where(DeviceEvent.device_id == uuid.UUID(km)))
            ).scalars()
        )
        assert {"moved_site", "deactivated"} <= types
        actions = set((await s.execute(select(AuditLog.action).where(AuditLog.entity == "device"))).scalars())
        assert {"update", "bulk_update", "bulk_deactivate"} <= actions

    # Técnico não edita; cliente visualizador também não.
    _, email = await factory.user(tenant.reseller_id, role="customer_viewer", customer_id=tenant.customer_id)
    viewer = await login(client, email)
    assert (
        await client.patch(f"/api/v1/devices/{km}", json={"sector": "x"}, headers=auth(viewer))
    ).status_code == 403


async def test_export_respects_filters(client: httpx.AsyncClient, factory: Factory) -> None:
    _, admin, _ = await setup(client, factory)
    resp = await client.get("/api/v1/park/export", params={"format": "xlsx", "q": "KM"}, headers=auth(admin))
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][6] == "Serial"
    assert [r[6] for r in rows[1:]] == ["KM-001"]
    assert rows[1][-1] == "C:55% M:66% Y:5% K:25%"
    csv = await client.get("/api/v1/park/export", params={"format": "csv"}, headers=auth(admin))
    text = csv.content.decode("utf-8-sig")
    assert text.count("\r\n") == 4  # cabeçalho + 3


async def test_counters_supplies_history_and_adjustments(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    tenant, admin, agent = await setup(client, factory)
    # Mais leituras em dias anteriores para a série diária.
    await agent.send(
        [
            agent.reading(
                "KM-001",
                {"total": 216000, "mono": 99500, "color": 116500},
                ip="10.0.0.11",
                read_at=sp_noon(2),
            ),
            agent.reading(
                "KM-001",
                {"total": 216500, "mono": 99800, "color": 116700},
                ip="10.0.0.11",
                read_at=sp_noon(1),
            ),
        ]
    )
    km = next(r for r in (await park(client, admin))["items"] if r["serial"] == "KM-001")["id"]
    series = (
        await client.get(f"/api/v1/devices/{km}/counters", params={"periods": 7}, headers=auth(admin))
    ).json()
    assert [p["total"] for p in series][-3:] == [216000, 216500, 217031]
    assert series[-1]["pages"] == 531
    assert series[-1]["pages_mono"] == 350
    assert series[-1]["pages_color"] == 181
    monthly = (
        await client.get(
            f"/api/v1/devices/{km}/counters",
            params={"granularity": "month", "periods": 2},
            headers=auth(admin),
        )
    ).json()
    assert monthly[-1]["total"] == 217031
    hist = (await client.get(f"/api/v1/devices/{km}/supplies/history", headers=auth(admin))).json()
    assert {h["color"] for h in hist} == {"black", "cyan", "magenta", "yellow"}

    exp = await client.get(
        f"/api/v1/devices/{km}/readings/export", params={"format": "csv"}, headers=auth(admin)
    )
    assert exp.status_code == 200
    assert exp.content.decode("utf-8-sig").count(CRLF) == 4  # cabeçalho + 3 leituras
    readings = (await client.get(f"/api/v1/devices/{km}/readings", headers=auth(admin))).json()["items"]
    last = readings[0]
    body = {
        "reading_id": last["id"],
        "read_at": last["read_at"],
        "total": 217000,
        "reason": "Contador lido errado na troca da placa",
    }
    resp = await client.post(f"/api/v1/devices/{km}/adjustments", json=body, headers=auth(admin))
    assert resp.status_code == 201, resp.text
    adj = resp.json()
    assert adj[0]["total"] == 217000
    assert adj[0]["user_name"] == "Usuário de Teste"
    empty = {**body, "total": None}
    assert (
        await client.post(f"/api/v1/devices/{km}/adjustments", json=empty, headers=auth(admin))
    ).status_code == 400
    short = {**body, "reason": "x"}
    assert (
        await client.post(f"/api/v1/devices/{km}/adjustments", json=short, headers=auth(admin))
    ).status_code == 422
    wrong = {**body, "reading_id": str(uuid.uuid4())}
    assert (
        await client.post(f"/api/v1/devices/{km}/adjustments", json=wrong, headers=auth(admin))
    ).status_code == 404
    async with sessionmaker() as s:
        # R6: a leitura original continua intacta.
        orig = (
            await s.execute(select(Reading.total).where(Reading.id == uuid.UUID(last["id"])))
        ).scalar_one()
        assert orig == 217031
        assert (await s.execute(select(ReadingAdjustment))).scalars().one().reason.startswith("Contador")
    _, email = await factory.user(tenant.reseller_id, role="technician")
    tech = await login(client, email)
    assert (
        await client.post(f"/api/v1/devices/{km}/adjustments", json=body, headers=auth(tech))
    ).status_code == 403


async def test_dashboard(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    tenant, admin, agent = await setup(client, factory)
    now = datetime.now(UTC)
    await agent.send(
        [
            agent.reading(
                "KM-001",
                {"total": 216000, "mono": 99500, "color": 116500},
                ip="10.0.0.11",
                read_at=sp_noon(1),
            )
        ]
    )
    d = (await client.get("/api/v1/dashboard", headers=auth(admin))).json()
    assert d["cards"]["devices_monitored"] == 3
    assert d["cards"]["devices_online"] == 3
    assert d["cards"]["devices_disconnected"] == 0
    assert d["cards"]["agents_online"] == 1
    assert d["cards"]["agents_offline"] == 0
    assert d["cards"]["toners_critical"] == 1  # amarelo 5%
    assert d["critical_supplies"][0]["color"] == "yellow"
    assert len(d["pages_per_day"]) == 30
    assert sum(p["mono"] for p in d["pages_per_day"]) == 650
    assert sum(p["color"] for p in d["pages_per_day"]) == 381
    # Coletor sem sinal há muito tempo aparece na lista "offline agora".
    async with sessionmaker() as s:
        await s.execute(update(Device).values(last_read_at=now - timedelta(hours=10)))
        await s.execute(update(Agent).values(state="offline", last_seen_at=now - timedelta(hours=1)))
        await s.commit()
        newly, back = await park_svc.mark_disconnected(s, 6)
        await s.commit()
    assert (newly, back) == (3, 0)
    d = (await client.get("/api/v1/dashboard", headers=auth(admin))).json()
    assert d["cards"]["devices_disconnected"] == 3
    assert d["cards"]["agents_offline"] == 1
    assert d["offline_agents"][0]["name"] == "Coletor 1"
    assert (await park(client, admin, disconnected=True))["total"] == 3
    # Leitura nova limpa o "desconectado".
    await agent.send(
        [agent.reading("KM-001", {"total": 217100, "mono": 100200, "color": 116900}, ip="10.0.0.11")]
    )
    assert (await park(client, admin, disconnected=True))["total"] == 2
    other = await factory.tenant("Revenda B")
    other_admin = await login(client, other.admin_email)
    d = (await client.get("/api/v1/dashboard", headers=auth(other_admin))).json()
    assert d["cards"]["devices_monitored"] == 0
    assert sum(p["mono"] for p in d["pages_per_day"]) == 0
    _ = tenant
