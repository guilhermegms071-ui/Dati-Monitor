"""Relatórios (PROMPT 10.9 / 16.12 / 15 itens 10-11): production without regressed counters, manual
adjustments, cutoff reading, monthly billing with franchise/overage, daily counter, the park/agent lists,
scope and the CSV/XLSX/PDF exports."""

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

import httpx
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from tests.agent_helpers import FakeAgent, enrolled_agent
from tests.alert_helpers import set_agent, set_device, supply
from tests.conftest import Factory, Tenant, auth, login

SERIAL = "RPT0001"


def month_before(today: date) -> date:
    return (today.replace(day=1) - timedelta(days=1)).replace(day=1)


def at(day: date, hour: int = 15) -> datetime:
    return datetime(day.year, day.month, day.day, hour, tzinfo=UTC)


class Park:
    """Previous month M: readings on M-1/20, M/10, M/15 (regression), M/20; replacement of black toner."""

    def __init__(self, tenant: Tenant, admin: str, agent: FakeAgent, month: date) -> None:
        self.tenant, self.admin, self.agent, self.month = tenant, admin, agent, month
        self.prev = month_before(month)
        self.last_day = (month.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)

    def day(self, n: int) -> date:
        return self.month.replace(day=n)


async def build_park(client: httpx.AsyncClient, factory: Factory) -> Park:
    tenant = await factory.tenant()
    admin = await login(client, tenant.admin_email)
    agent = await enrolled_agent(client, tenant)
    today = datetime.now(UTC).date()
    park = Park(tenant, admin, agent, month_before(today))
    points = [
        (at(park.prev.replace(day=20)), {"total": 1000, "mono": 800, "color": 200}),
        (at(park.day(10)), {"total": 1500, "mono": 1100, "color": 400}),
        (at(park.day(15)), {"total": 900, "mono": 700, "color": 200}),  # regressão: fica fora
        (at(park.day(20)), {"total": 1600, "mono": 1150, "color": 450}),
    ]
    for read_at, counters in points:
        results = await agent.send([agent.reading(SERIAL, counters, read_at=read_at)])
        assert results[0]["status"] == "accepted", results
    # Troca de toner preto (5% → 95%) com contadores antes/depois.
    for read_at, pct in ((at(park.day(10), 16), 5.0), (at(park.day(20), 16), 95.0)):
        item = agent.item("supplies", serial=SERIAL, read_at=read_at, supplies=[supply("k", "black", pct)])
        assert (await agent.send([item]))[0]["status"] == "accepted"
    await agent.send([agent.reading("RPT0002", {"total": 50, "mono": 50, "color": 0}, ip="10.0.0.6")])
    return park


async def report(client: httpx.AsyncClient, token: str, key: str, **params: Any) -> dict[str, Any]:
    resp = await client.get(f"/api/v1/reports/{key}", params=params, headers=auth(token))
    assert resp.status_code == 200, resp.text
    data: dict[str, Any] = resp.json()
    return data


def by_serial(data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {r["serial"]: r for r in data["rows"]}


async def test_catalog(client: httpx.AsyncClient, factory: Factory) -> None:
    tenant = await factory.tenant()
    admin = await login(client, tenant.admin_email)
    catalog = (await client.get("/api/v1/reports", headers=auth(admin))).json()
    keys = {r["key"] for r in catalog}
    assert {
        "production", "daily_counter", "cutoff", "billing", "supply_replacements", "supply_yield",
        "supply_alerts", "park_overview", "park_status", "online", "disconnected", "deactivated",
        "no_reading", "discoveries", "ip_changes", "regressions", "agent_status", "agent_availability",
    } <= keys  # fmt: skip
    cutoff = next(r for r in catalog if r["key"] == "cutoff")
    assert (cutoff["period"], cutoff["cutoff_date"]) == (False, True)
    billing = next(r for r in catalog if r["key"] == "billing")
    assert billing["month"] is True
    assert (await client.get("/api/v1/reports/nao-existe", headers=auth(admin))).status_code == 404


async def test_production_skips_regressions_and_uses_adjustments(
    client: httpx.AsyncClient, factory: Factory
) -> None:
    park = await build_park(client, factory)
    period = {"date_from": park.month.isoformat(), "date_to": park.last_day.isoformat()}
    prod = by_serial(await report(client, park.admin, "production", **period))
    # Base 1000 (mês anterior) → 1500 → [900 regressão ignorada] → 1600.
    assert (prod[SERIAL]["total"], prod[SERIAL]["mono"], prod[SERIAL]["color"]) == (600, 350, 250)
    assert prod[SERIAL]["readings"] == 2

    # Ajuste manual na última leitura: a produção passa a usar o valor corrigido.
    device_id = await set_device(factory.maker, SERIAL)
    readings = (await client.get(f"/api/v1/devices/{device_id}/readings", headers=auth(park.admin))).json()[
        "items"
    ]
    last = next(r for r in readings if r["total"] == 1600)
    body = {"reading_id": last["id"], "read_at": last["read_at"], "total": 1700, "mono": 1200, "color": 500,
            "reason": "Conferido na folha de contadores"}  # fmt: skip
    assert (
        await client.post(f"/api/v1/devices/{device_id}/adjustments", json=body, headers=auth(park.admin))
    ).status_code == 201
    prod = by_serial(await report(client, park.admin, "production", **period))
    assert (prod[SERIAL]["total"], prod[SERIAL]["mono"], prod[SERIAL]["color"]) == (700, 400, 300)

    by_customer = await report(client, park.admin, "production", group_by="customer", **period)
    assert by_customer["rows"] == [
        {"customer": "Revenda A Cliente", "devices": 2, "mono": 400, "color": 300, "total": 700}
    ]
    assert by_customer["totals"]["total"] == 700

    daily = await report(client, park.admin, "daily_counter", **period)
    days = {r["day"]: r for r in daily["chart_rows"]}
    assert days[park.day(10).isoformat()]["mono"] == 300
    assert days[park.day(20).isoformat()]["color"] == 100
    assert days[park.day(15).isoformat()]["mono"] == 0
    assert daily["totals"]["total"] == 700
    assert daily["chart"]["stacked"] is True
    assert daily["total_rows"] == park.last_day.day

    regressions = await report(client, park.admin, "regressions", **period)
    received = await report(client, park.admin, "regressions", date_type="received", **period)
    assert received["rows"] == [], "recebidas hoje, fora do mês passado"
    assert {(r["counter"], r["before"], r["after"], r["review"]) for r in regressions["rows"]} == {
        ("Total", 1500, 900, "Não classificada"),
        ("PB", 1100, 700, "Não classificada"),
        ("Cor", 400, 200, "Não classificada"),
    }


async def test_cutoff_and_billing(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    park = await build_park(client, factory)
    cut = by_serial(await report(client, park.admin, "cutoff", date_to=park.day(16).isoformat()))
    # A leitura do dia 15 é regressão: a de corte continua sendo a do dia 10.
    assert (cut[SERIAL]["total"], cut[SERIAL]["mono"]) == (1500, 1100)
    assert cut[SERIAL]["read_at"].startswith(park.day(10).isoformat())
    assert cut[SERIAL]["erp_code"] == "E-Revenda A"
    assert cut["RPT0002"]["read_at"] is None or cut["RPT0002"]["read_at"] > park.day(16).isoformat()

    await set_device(
        sessionmaker,
        SERIAL,
        franchise_value=Decimal("100.00"),
        franchise_pages_mono=300,
        franchise_pages_color=100,
        overage_price_mono=Decimal("0.05"),
        overage_price_color=Decimal("0.40"),
    )
    bill = by_serial(await report(client, park.admin, "billing", month=park.month.strftime("%Y-%m")))
    row = bill[SERIAL]
    assert (row["start_mono"], row["end_mono"]) == (800, 1150)
    assert (row["pages_mono"], row["pages_color"]) == (350, 250)
    assert (row["excess_mono"], row["excess_color"]) == (50, 150)
    assert row["overage_value"] == 62.5  # 50 x 0,05 + 150 x 0,40
    assert row["billed"] == 162.5
    assert bill["RPT0002"]["billed"] is None  # sem franquia nem preço cadastrados


async def test_park_lists_supplies_and_agents(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    park = await build_park(client, factory)
    overview = by_serial(await report(client, park.admin, "park_overview"))
    assert set(overview) == {SERIAL, "RPT0002"}
    assert overview[SERIAL]["levels"] == "K 95%"
    await set_device(sessionmaker, "RPT0002", active=False)
    assert set(by_serial(await report(client, park.admin, "deactivated"))) == {"RPT0002"}
    assert set(by_serial(await report(client, park.admin, "online"))) == {SERIAL}
    status = (await report(client, park.admin, "park_status"))["rows"]
    assert (status[0]["devices"], status[0]["inactive"], status[0]["online"]) == (2, 1, 1)
    # Sem leitura há 24 h: a última leitura do RPT0001 é do mês passado.
    assert set(by_serial(await report(client, park.admin, "no_reading", hours=24))) == {SERIAL}

    period = {"date_from": park.month.isoformat(), "date_to": park.last_day.isoformat()}
    repl = (await report(client, park.admin, "supply_replacements", **period))["rows"]
    assert len(repl) == 1
    assert (repl[0]["color"], repl[0]["level_before"], repl[0]["level_after"]) == ("Preto", 5.0, 95.0)
    yld = (await report(client, park.admin, "supply_yield", **period))["rows"]
    assert (yld[0]["count"], yld[0]["model"]) == (1, "Konica Minolta bizhub C287")

    agents = (await report(client, park.admin, "agent_status"))["rows"]
    assert agents[0]["agent"] == "Coletor 1"
    assert agents[0]["devices"] == 1  # só os ativos
    await set_agent(sessionmaker, park.agent.agent_id, enrolled_at=datetime.now(UTC) - timedelta(minutes=30))
    await park.agent.heartbeat()
    avail = (await report(client, park.admin, "agent_availability"))["rows"]
    assert avail[0]["online_pct"] is not None
    assert 0 < avail[0]["online_pct"] <= 100
    by_customer = (await report(client, park.admin, "agent_availability", group_by="customer"))["rows"]
    assert by_customer[0]["agents"] == 1


async def test_scope_validation_and_exports(client: httpx.AsyncClient, factory: Factory) -> None:
    park = await build_park(client, factory)
    other = await factory.tenant("Outra empresa")
    stranger = await login(client, other.admin_email)
    assert not (await report(client, stranger, "park_overview"))["rows"]
    resp = await client.get(
        "/api/v1/reports/park_overview",
        params={"customer_id": str(park.tenant.customer_id)},
        headers=auth(stranger),
    )
    assert resp.status_code == 404
    bad = await client.get(
        "/api/v1/reports/production", params={"date_from": "2026-05-10", "date_to": "2026-05-01"},
        headers=auth(park.admin),
    )  # fmt: skip
    assert bad.json()["detail"]["code"] == "invalid_period"
    bad = await client.get(
        "/api/v1/reports/production", params={"group_by": "planeta"}, headers=auth(park.admin)
    )
    assert bad.json()["detail"]["code"] == "invalid_group_by"

    period = {"date_from": park.month.isoformat(), "date_to": park.last_day.isoformat()}
    url = "/api/v1/reports/production/export"
    csv_resp = await client.get(url, params={**period, "format": "csv"}, headers=auth(park.admin))
    assert csv_resp.status_code == 200
    text = csv_resp.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Cliente;Local;Serial")
    assert "Total;" in text.splitlines()[-1]
    xlsx = await client.get(url, params={**period, "format": "xlsx"}, headers=auth(park.admin))
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert ws is not None
    assert [c.value for c in ws[1]][:3] == ["Cliente", "Local", "Serial"]
    pdf = await client.get(url, params={**period, "format": "pdf"}, headers=auth(park.admin))
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-disposition"].endswith('.pdf"')
    # Relatório sem linhas também gera PDF (com o aviso).
    empty = await client.get(
        "/api/v1/reports/regressions/export",
        params={"format": "pdf", "date_from": "2026-01-01", "date_to": "2026-01-02"},
        headers=auth(park.admin),
    )
    assert empty.status_code == 200


async def test_device_counters_period_and_daily(client: httpx.AsyncClient, factory: Factory) -> None:
    """Contadores por equipamento (formato do Datacount): um bloco por equipamento com leitura inicial,
    final e diferença; ou dia a dia, com o contador do dia e as páginas do dia."""
    park = await build_park(client, factory)
    period = {
        "date_from": park.month.isoformat(),
        "date_to": park.last_day.isoformat(),
        "customer_id": str(park.tenant.customer_id),
    }
    data = await report(client, park.admin, "device_counters", **period)
    assert data["filters"]["group_by"] == "period"
    rows = [r for r in data["rows"] if r["serial"] == SERIAL]
    lines = {r["info"]: (r["first"], r["last"], r["diff"]) for r in rows}
    # Base 1000 (mês anterior) → 1600; a regressão do dia 15 fica fora.
    assert lines["Total geral"] == (1000, 1600, 600)
    assert lines["Total PB"] == (800, 1150, 350)
    assert lines["Total Cor"] == (200, 450, 250)
    sections = {s["key"]: s for s in data["sections"]}
    section = sections[rows[0]["_section"]]
    assert SERIAL in section["title"]
    assert {d["label"]: d["value"] for d in section["details"]}["Leitura final"].startswith(
        park.day(20).strftime("%d/%m/%Y")
    )
    assert {c["key"] for c in data["columns"] if c["section"]} == {
        "customer",
        "site",
        "serial",
        "model",
        "sector",
    }
    summary = {s["label"]: s["value"] for s in data["summary"]}
    assert (summary["Equipamentos"], summary["Total de páginas"], summary["Com leitura no período"]) == (
        2,
        600,
        1,
    )

    daily = await report(client, park.admin, "device_counters", group_by="daily", limit=500, **period)
    days = {r["day"]: r for r in daily["rows"] if r["serial"] == SERIAL}
    assert len(days) == park.last_day.day
    d10, d15, d20 = (days[park.day(n).isoformat()] for n in (10, 15, 20))
    assert (d10["counter_total"], d10["pages_total"], d10["pages_mono"]) == (1500, 500, 300)
    assert (d15["counter_total"], d15["pages_total"]) == (None, 0)  # regressão não conta
    assert (d20["counter_total"], d20["pages_total"]) == (1600, 100)
    sec = {s["key"]: s for s in daily["sections"]}[d10["_section"]]
    assert sec["totals"]["pages_total"] == 600
    assert daily["totals"]["pages_total"] == 600
    chart = {r["day"]: r for r in daily["chart_rows"]}
    assert len(chart) == park.last_day.day
    assert chart[park.day(10).isoformat()]["mono"] == 300

    read_days = await report(client, park.admin, "device_counters", group_by="daily_read", **period)
    assert sorted(r["day"] for r in read_days["rows"]) == [park.day(10).isoformat(), park.day(20).isoformat()]
    assert len(read_days["chart_rows"]) == park.last_day.day
    assert {r["read_at"] for r in read_days["rows"]} == {"12:00"}  # 15:00 UTC = 12:00 em Brasília

    for view in ("period", "daily", "daily_read"):
        for fmt in ("pdf", "xlsx", "csv"):
            resp = await client.get(
                "/api/v1/reports/device_counters/export",
                params={**period, "group_by": view, "format": fmt},
                headers=auth(park.admin),
            )
            assert resp.status_code == 200, resp.text
    csv_text = resp.content.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("Cliente;Local;Nº de série;Modelo;Setor;Dia")
