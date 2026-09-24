"""Unit tests for pure helpers: crypto, tokens, rate limiter, validators, cursors, exports."""

import io
import uuid
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core import crypto
from app.core.errors import AppError
from app.core.permissions import ALL_PERMISSIONS, ROLE_PERMISSIONS, ROLES
from app.core.ratelimit import RateLimiter
from app.core.security import (
    InvalidTokenError,
    create_access_token,
    decode_access_token,
    hash_password,
    hash_token,
    verify_password,
)
from app.core.validators import check_password_policy, check_timezone, normalize_cnpj, normalize_email
from app.services.audit import diff
from app.services.export import ExportColumn, to_csv, to_xlsx
from app.services.pagination import decode_cursor, encode_cursor

KEY = b"0" * 32


def test_crypto_roundtrip_and_tamper_detection() -> None:
    blob = crypto.encrypt(KEY, "comunidade-secreta", associated_data=b"site:1")
    assert b"comunidade" not in blob
    assert crypto.decrypt(KEY, blob, associated_data=b"site:1") == "comunidade-secreta"
    with pytest.raises(crypto.DecryptionError):
        crypto.decrypt(KEY, blob, associated_data=b"site:2")
    with pytest.raises(crypto.DecryptionError):
        crypto.decrypt(b"1" * 32, blob, associated_data=b"site:1")
    with pytest.raises(crypto.DecryptionError):
        crypto.decrypt(KEY, b"\x02" + blob[1:], associated_data=b"site:1")
    data = crypto.encrypt_json(KEY, {"token": "abc"})
    assert crypto.decrypt_json(KEY, data) == {"token": "abc"}
    with pytest.raises(crypto.DecryptionError):
        crypto.decrypt_json(KEY, crypto.encrypt(KEY, "[1, 2]"))


def test_password_hashing() -> None:
    h = hash_password("Senha-Forte-123")
    assert h.startswith("$argon2id$")
    assert verify_password(h, "Senha-Forte-123")
    assert not verify_password(h, "errada")
    assert not verify_password(None, "qualquer")
    assert not verify_password("não-é-hash", "x")
    assert hash_token("a") == hash_token("a")
    assert hash_token("a") != hash_token("b")


def test_access_token_roundtrip_and_errors() -> None:
    uid, rid, cid = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    changed = datetime(2026, 1, 1, tzinfo=UTC)
    token, exp = create_access_token(
        secret="s" * 40,
        user_id=uid,
        reseller_id=rid,
        role="operator",
        customer_id=cid,
        minutes=15,
        limited="password_change_required",
        password_changed_at=changed,
    )
    claims = decode_access_token(token, secret="s" * 40)
    assert (claims.user_id, claims.reseller_id, claims.customer_id) == (uid, rid, cid)
    assert claims.limited == "password_change_required"
    assert claims.expires_at == exp
    assert claims.password_version == int(changed.timestamp() * 1_000_000)
    with pytest.raises(InvalidTokenError, match="inválido"):
        decode_access_token(token, secret="t" * 40)
    old, _ = create_access_token(
        secret="s" * 40,
        user_id=uid,
        reseller_id=rid,
        role="operator",
        customer_id=None,
        minutes=1,
        now=datetime.now(UTC) - timedelta(hours=1),
    )
    with pytest.raises(InvalidTokenError, match="expirado"):
        decode_access_token(old, secret="s" * 40)
    import jwt  # noqa: PLC0415

    refresh_like = jwt.encode(
        {"typ": "refresh", "sub": str(uid), "iat": 1, "exp": 9999999999}, "s" * 40, "HS256"
    )
    with pytest.raises(InvalidTokenError, match="tipo"):
        decode_access_token(refresh_like, secret="s" * 40)
    broken = jwt.encode({"typ": "access", "sub": "x", "iat": 1, "exp": 9999999999}, "s" * 40, "HS256")
    with pytest.raises(InvalidTokenError, match="dados inválidos"):
        decode_access_token(broken, secret="s" * 40)


def test_rate_limiter_window() -> None:
    rl = RateLimiter(2, 10)
    assert rl.hit("a", now=0)
    assert rl.hit("a", now=1)
    assert not rl.hit("a", now=2)
    assert rl.retry_after("a", now=2) == pytest.approx(8)
    assert rl.hit("b", now=2)
    assert rl.hit("a", now=10.5)
    assert rl.retry_after("zzz") == 0.0


def test_validators() -> None:
    assert normalize_cnpj("11.222.333/0001-81") == "11222333000181"
    assert normalize_cnpj("  ") is None
    for bad in ("11.222.333/0001-82", "00000000000000", "123"):
        with pytest.raises(ValueError, match="CNPJ"):
            normalize_cnpj(bad)
    assert normalize_email(" Fulano@Empresa.COM ") == "fulano@empresa.com"
    with pytest.raises(ValueError, match="e-mail"):
        normalize_email("sem-arroba")
    assert check_password_policy("Senha-Forte-1") == "Senha-Forte-1"
    for bad_pw, msg in (("curta", "10"), ("a" * 12, "fraca"), ("x" * 300, "256")):
        with pytest.raises(ValueError, match=msg):
            check_password_policy(bad_pw)
    with pytest.raises(ValueError, match="e-mail"):
        check_password_policy("fulano@empresa.com", "fulano@empresa.com")
    assert check_timezone("America/Sao_Paulo") == "America/Sao_Paulo"
    with pytest.raises(ValueError, match="fuso"):
        check_timezone("Terra/Plana")


def test_permission_matrix_is_consistent() -> None:
    assert {r.code for r in ROLES} == set(ROLE_PERMISSIONS)
    for perms in ROLE_PERMISSIONS.values():
        assert perms <= ALL_PERMISSIONS
    assert ROLE_PERMISSIONS["superadmin"] == ALL_PERMISSIONS
    assert "resellers.write" not in ROLE_PERMISSIONS["reseller_admin"]
    assert not any(p.endswith(".write") for p in ROLE_PERMISSIONS["customer_viewer"])


def test_cursor_roundtrip_and_validation() -> None:
    rid = uuid.uuid4()
    now = datetime(2026, 9, 24, 12, 0, tzinfo=UTC)
    c = encode_cursor("created_at", "desc", now, rid)
    assert decode_cursor(c, "created_at", "desc", "datetime") == (now, rid)
    c2 = encode_cursor("value", "asc", Decimal("1.50"), rid)
    assert decode_cursor(c2, "value", "asc", "decimal") == (Decimal("1.50"), rid)
    assert decode_cursor(encode_cursor("n", "asc", 5, rid), "n", "asc", "int") == (5, rid)
    with pytest.raises(AppError):
        decode_cursor(c, "created_at", "asc", "datetime")
    with pytest.raises(AppError):
        decode_cursor("!!!", "created_at", "desc", "datetime")


@dataclass
class _Row:
    name: str
    when: datetime | None
    ok: bool
    tags: list[str]
    amount: Decimal


def test_csv_and_xlsx_export_formatting() -> None:
    rows = [
        _Row("Ação", datetime(2026, 1, 1, 15, 0, tzinfo=UTC), True, ["a", "b"], Decimal("2.5")),
        _Row("B", None, False, [], Decimal(0)),
    ]
    cols = [
        ExportColumn[_Row]("Nome", lambda r: r.name),
        ExportColumn[_Row]("Quando", lambda r: r.when),
        ExportColumn[_Row]("OK", lambda r: r.ok),
        ExportColumn[_Row]("Tags", lambda r: r.tags),
        ExportColumn[_Row]("Valor", lambda r: r.amount),
    ]
    csv_text = to_csv(rows, cols).decode("utf-8")
    assert csv_text.startswith("﻿Nome;Quando;OK;Tags;Valor")
    # 15:00 UTC = 12:00 em São Paulo
    assert "Ação;01/01/2026 12:00:00;Sim;a, b;2.5" in csv_text
    assert "B;;Não;;0" in csv_text
    wb = load_workbook(io.BytesIO(to_xlsx(rows, cols, "Planilha com nome muito longo para o Excel aceitar")))
    ws = wb.active
    assert ws is not None
    assert ws.title == "Planilha com nome muito longo p"
    values = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert values[0] == "Ação"
    assert values[1] == datetime(2026, 1, 1, 12, 0)


def test_audit_diff_ignores_unchanged_and_updated_at() -> None:
    before = {"a": 1, "b": 2, "updated_at": "x"}
    after = {"a": 1, "b": 3, "updated_at": "y"}
    assert diff(before, after) == ({"b": 2}, {"b": 3})
