"""CRUD of resellers, companies, customers and sites through the API (with audit, pagination and export)."""

import io
import uuid

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.models import AuditLog, Customer, Device
from tests.conftest import Factory, auth, login

pytestmark = pytest.mark.usefixtures("clean_db")


async def test_full_hierarchy_crud_with_audit(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    t = await factory.tenant()
    h = auth(await login(client, t.admin_email))

    company = await client.post(
        "/api/v1/companies", json={"legal_name": "Nova Empresa", "cnpj": "11.222.333/0001-81"}, headers=h
    )
    assert company.status_code == 201, company.text
    cid = company.json()["id"]
    assert company.json()["cnpj"] == "11222333000181"

    bad_cnpj = await client.post(
        "/api/v1/companies", json={"legal_name": "X Y", "cnpj": "11.111.111/1111-11"}, headers=h
    )
    assert bad_cnpj.status_code == 422

    customer = await client.post(
        "/api/v1/customers",
        json={
            "company_id": cid,
            "name": "Hospital Central",
            "erp_code": "ERP-77",
            "email": "TI@Hospital.test",
        },
        headers=h,
    )
    assert customer.status_code == 201, customer.text
    cust = customer.json()
    assert cust["email"] == "ti@hospital.test"

    dup = await client.post(
        "/api/v1/customers", json={"company_id": cid, "name": "Outro", "erp_code": "ERP-77"}, headers=h
    )
    assert dup.status_code == 409
    assert dup.json()["detail"]["code"] == "erp_code_taken"

    site = await client.post(
        "/api/v1/sites",
        json={"customer_id": cust["id"], "name": "Bloco A", "timezone": "America/Sao_Paulo"},
        headers=h,
    )
    assert site.status_code == 201, site.text
    bad_tz = await client.post(
        "/api/v1/sites", json={"customer_id": cust["id"], "name": "B", "timezone": "Marte/Base"}, headers=h
    )
    assert bad_tz.status_code == 422

    upd = await client.patch(
        f"/api/v1/customers/{cust['id']}",
        json={"name": "Hospital Central II", "phone": "21 99999-0000"},
        headers=h,
    )
    assert upd.status_code == 200
    assert upd.json()["name"] == "Hospital Central II"
    null_name = await client.patch(f"/api/v1/customers/{cust['id']}", json={"name": None}, headers=h)
    assert null_name.status_code == 400

    site_upd = await client.patch(
        f"/api/v1/sites/{site.json()['id']}",
        json={"address": "Rua 1", "collection_config": {"counters_minutes": 30}},
        headers=h,
    )
    assert site_upd.json()["collection_config"] == {"counters_minutes": 30}

    in_use = await client.delete(f"/api/v1/customers/{cust['id']}", headers=h)
    assert in_use.status_code == 409
    assert (await client.delete(f"/api/v1/sites/{site.json()['id']}", headers=h)).status_code == 204
    assert (await client.delete(f"/api/v1/companies/{cid}", headers=h)).status_code == 409
    assert (await client.delete(f"/api/v1/customers/{cust['id']}", headers=h)).status_code == 204
    assert (await client.get(f"/api/v1/customers/{cust['id']}", headers=h)).status_code == 404
    assert (await client.delete(f"/api/v1/companies/{cid}", headers=h)).status_code == 204

    async with sessionmaker() as s:
        rows = (await s.execute(select(AuditLog).where(AuditLog.entity == "customer"))).scalars().all()
    actions = sorted(r.action for r in rows)
    assert actions == ["create", "delete", "update"]
    update_row = next(r for r in rows if r.action == "update")
    assert update_row.before == {"name": "Hospital Central", "phone": None}
    assert update_row.after == {"name": "Hospital Central II", "phone": "21 99999-0000"}
    assert update_row.ip == "203.0.113.10"


async def test_customer_pagination_sort_and_filters(client: httpx.AsyncClient, factory: Factory) -> None:
    t = await factory.tenant()
    h = auth(await login(client, t.admin_email))
    for i in range(7):
        r = await client.post(
            "/api/v1/customers",
            json={"company_id": str(t.company_id), "name": f"Cliente {i:02d}", "active": i % 2 == 0},
            headers=h,
        )
        assert r.status_code == 201
    seen: list[str] = []
    cursor = None
    while True:
        params: dict[str, str | int] = {"limit": 3, "sort": "name", "direction": "desc"}
        if cursor:
            params["cursor"] = cursor
        page = (await client.get("/api/v1/customers", params=params, headers=h)).json()
        seen += [c["name"] for c in page["items"]]
        cursor = page["next_cursor"]
        if not cursor:
            break
    assert seen == sorted(seen, reverse=True)
    assert len(seen) == 8  # 7 novos + o do tenant
    inactive = (await client.get("/api/v1/customers", params={"active": "false"}, headers=h)).json()["items"]
    assert {c["name"] for c in inactive} == {"Cliente 01", "Cliente 03", "Cliente 05"}
    q = (await client.get("/api/v1/customers", params={"q": "cliente 0%"}, headers=h)).json()["items"]
    assert q == []  # "%" é literal, não curinga
    bad_sort = await client.get("/api/v1/customers", params={"sort": "senha"}, headers=h)
    assert bad_sort.status_code == 400
    bad_cursor = await client.get("/api/v1/customers", params={"cursor": "xxx"}, headers=h)
    assert bad_cursor.json()["detail"]["code"] == "invalid_cursor"
    wrong_sort_cursor = (await client.get("/api/v1/customers", params={"limit": 1}, headers=h)).json()[
        "next_cursor"
    ]
    mismatch = await client.get(
        "/api/v1/customers", params={"cursor": wrong_sort_cursor, "sort": "created_at"}, headers=h
    )
    assert mismatch.status_code == 400


async def test_customer_export_csv_and_xlsx_respect_filters(
    client: httpx.AsyncClient, factory: Factory
) -> None:
    t = await factory.tenant()
    h = auth(await login(client, t.admin_email))
    await client.post(
        "/api/v1/customers",
        json={"company_id": str(t.company_id), "name": "Ação Ltda", "erp_code": "Z1"},
        headers=h,
    )
    csv_resp = await client.get("/api/v1/customers/export", params={"format": "csv", "q": "Ação"}, headers=h)
    assert csv_resp.status_code == 200
    assert csv_resp.headers["content-type"].startswith("text/csv")
    text = csv_resp.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0].startswith("Nome;CNPJ;Código ERP")
    assert len(lines) == 2
    assert "Ação Ltda;;Z1" in lines[1]
    xlsx = await client.get("/api/v1/customers/export", params={"format": "xlsx"}, headers=h)
    wb = load_workbook(io.BytesIO(xlsx.content))
    ws = wb.active
    assert ws is not None
    names = sorted(str(row[0]) for row in ws.iter_rows(min_row=2, values_only=True))
    assert names == ["Ação Ltda", "Revenda A Cliente"]
    assert 'filename="clientes-' in xlsx.headers["content-disposition"]


async def test_superadmin_manages_resellers(client: httpx.AsyncClient, factory: Factory) -> None:
    t = await factory.tenant(admin_role="superadmin")
    h = auth(await login(client, t.admin_email))
    created = await client.post("/api/v1/resellers", json={"name": "Revenda Nova"}, headers=h)
    assert created.status_code == 201
    rid = created.json()["id"]
    listed = (await client.get("/api/v1/resellers", headers=h)).json()["items"]
    assert {r["name"] for r in listed} == {"Revenda A", "Revenda Nova"}
    comp = await client.post(
        "/api/v1/companies", json={"legal_name": "Empresa da Nova", "reseller_id": rid}, headers=h
    )
    assert comp.json()["reseller_id"] == rid
    assert (await client.delete(f"/api/v1/resellers/{rid}", headers=h)).status_code == 409
    assert (await client.delete(f"/api/v1/companies/{comp.json()['id']}", headers=h)).status_code == 204
    assert (
        await client.patch(f"/api/v1/resellers/{rid}", json={"name": "Revenda Renomeada"}, headers=h)
    ).status_code == 200
    assert (await client.delete(f"/api/v1/resellers/{rid}", headers=h)).status_code == 204
    own = await client.delete(f"/api/v1/resellers/{t.reseller_id}", headers=h)
    assert own.json()["detail"]["code"] == "cannot_delete_own_reseller"


async def test_reseller_admin_cannot_write_resellers_but_reads_own(
    client: httpx.AsyncClient, factory: Factory
) -> None:
    t = await factory.tenant()
    h = auth(await login(client, t.admin_email))
    assert (await client.post("/api/v1/resellers", json={"name": "X Y"}, headers=h)).status_code == 403
    own = await client.get(f"/api/v1/resellers/{t.reseller_id}", headers=h)
    assert own.status_code == 200
    assert (await client.get("/api/v1/resellers", headers=h)).json()["items"][0]["id"] == str(t.reseller_id)
    other = await client.post(
        "/api/v1/companies", json={"legal_name": "Z Z", "reseller_id": str(uuid.uuid4())}, headers=h
    )
    assert other.status_code == 403


async def test_role_permissions_on_writes(client: httpx.AsyncClient, factory: Factory) -> None:
    t = await factory.tenant()
    _, tech = await factory.user(t.reseller_id, role="technician")
    _, oper = await factory.user(t.reseller_id, role="operator")
    ht = auth(await login(client, tech))
    ho = auth(await login(client, oper))
    body = {"company_id": str(t.company_id), "name": "Por Técnico"}
    assert (await client.post("/api/v1/customers", json=body, headers=ht)).status_code == 403
    assert (await client.get("/api/v1/customers", headers=ht)).status_code == 200
    assert (await client.post("/api/v1/customers", json=body, headers=ho)).status_code == 201
    assert (await client.get("/api/v1/users", headers=ho)).status_code == 403
    assert (await client.get("/api/v1/audit", headers=ho)).status_code == 403


async def test_site_delete_blocked_by_devices(
    client: httpx.AsyncClient, factory: Factory, sessionmaker: async_sessionmaker[AsyncSession]
) -> None:
    t = await factory.tenant()
    h = auth(await login(client, t.admin_email))
    async with sessionmaker() as s:
        s.add(Device(reseller_id=t.reseller_id, site_id=t.site_id, customer_id=t.customer_id, serial="SER-1"))
        await s.commit()
    resp = await client.delete(f"/api/v1/sites/{t.site_id}", headers=h)
    assert resp.status_code == 409
    assert "equipamentos" in resp.json()["detail"]["message"]
    async with sessionmaker() as s:
        assert (await s.get(Customer, t.customer_id)) is not None
