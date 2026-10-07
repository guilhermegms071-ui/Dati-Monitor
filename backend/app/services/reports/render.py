"""Report output: JSON cells for the screen, and CSV / XLSX / PDF files with the same columns."""

import io
from datetime import date, datetime
from decimal import Decimal
from importlib.resources import files
from typing import Any, Literal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.errors import bad_request
from app.services.export import DISPLAY_TZ, ExportColumn, export_value, to_csv
from app.services.reports.base import SECTION_KEY, Col, ReportData, Section, Stat

FileFormat = Literal["csv", "xlsx", "pdf"]
MAX_PDF_ROWS = 5000
MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def json_cell(value: Any) -> Any:
    """Decimals as numbers (the portal formats them), datetimes/dates as ISO."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def fmt_int(v: int | float) -> str:
    return f"{v:,.0f}".replace(",", ".")


def fmt_dec(v: Decimal | float, places: int = 1) -> str:
    s = f"{v:,.{places}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def text_cell(col: Col, value: Any) -> str:
    """How a value is printed in the PDF (pt-BR)."""
    if value is None or value == "":
        return "—"
    match col.kind:
        case "int":
            return fmt_int(value)
        case "money":
            return "R$ " + fmt_dec(value, 2)
        case "percent":
            return fmt_dec(value) + "%"
        case "decimal":
            return fmt_dec(value)
        case "bool":
            return "Sim" if value else "Não"
        case "datetime":
            return (
                value.astimezone(DISPLAY_TZ).strftime("%d/%m/%Y %H:%M")
                if isinstance(value, datetime)
                else str(value)
            )
        case "date":
            return value.strftime("%d/%m/%Y") if isinstance(value, date) else str(value)
    return str(value)


def _export_shape(data: ReportData) -> tuple[list[Col], list[dict[str, Any]], list[str]]:
    """Columns, rows (with the totals row) and CSV headers of the CSV/XLSX: the report's own export format
    when it has one (e.g. one line per device), else the same columns as the screen."""
    cols = data.export_columns or data.columns
    rows = data.export_rows if data.export_rows is not None else data.rows
    totals = data.export_totals if data.export_columns else data.totals
    if totals:
        rows = [*rows, {**totals, cols[0].key: "Total" if not data.export_columns else "Totais"}]
    headers = [c.label for c in cols]
    if data.export_groups:
        prefixes = [label for label, n in data.export_groups for _ in range(n)]
        headers = [f"{p}: {h}" if p else h for p, h in zip(prefixes, headers, strict=True)]
    return cols, rows, headers


def render(
    data: ReportData, fmt: FileFormat, *, title: str, meta: list[tuple[str, str]] | None = None
) -> bytes:
    """`meta` = filtros aplicados (rótulo, valor) mostrados no cabeçalho do PDF."""
    if fmt in ("csv", "xlsx"):
        cols, rows, headers = _export_shape(data)
        if fmt == "csv":
            return to_csv(rows, [ExportColumn(h, _getter(c.key)) for c, h in zip(cols, headers, strict=True)])
        return _xlsx(data, cols, rows, title)
    rows = data.rows
    if len(rows) > MAX_PDF_ROWS:
        raise bad_request(
            "pdf_too_large",
            f"O PDF aceita até {fmt_int(MAX_PDF_ROWS)} linhas; este relatório tem {fmt_int(len(rows))}. "
            "Use CSV ou XLSX, ou filtre por cliente.",
        )
    return _pdf(data, title=title, meta=meta or [])


def _getter(key: str) -> Any:
    return lambda row: row.get(key)


XLSX_FORMATS = {
    "int": "#,##0",
    "decimal": "#,##0.0",
    "money": '"R$" #,##0.00',
    "percent": '0.0"%"',
    "datetime": "dd/mm/yyyy hh:mm",
    "date": "dd/mm/yyyy",
}


def _xlsx(data: ReportData, cols: list[Col], rows: list[dict[str, Any]], title: str) -> bytes:
    """XLSX with styled header (and the group band when the report has one), number formats, widths,
    frozen header, filter and a bold totals row. Row 1 is always the header (or the group band)."""
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("planilha sem aba ativa")
    ws.title = title[:31]
    head_fill = PatternFill("solid", fgColor="156CC4")
    band_fill = PatternFill("solid", fgColor="143E6A")
    total_fill = PatternFill("solid", fgColor="E6EEF7")
    white = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="DBE2E8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    head_row = 1
    if data.export_groups:
        col = 1
        for label, n in data.export_groups:
            ws.cell(row=1, column=col, value=label)
            if n > 1:
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + n - 1)
            for ci in range(col, col + n):
                band = ws.cell(row=1, column=ci)
                band.fill, band.font, band.alignment, band.border = band_fill, white, center, border
            col += n
        head_row = 2
    for i, c in enumerate(cols, start=1):
        head = ws.cell(row=head_row, column=i, value=c.label)
        head.fill, head.font, head.alignment, head.border = head_fill, white, center, border
    totals_at = len(rows) - 1 if (data.export_totals if data.export_columns else data.totals) else -1
    for r_index, row in enumerate(rows):
        excel_row = head_row + 1 + r_index
        for i, c in enumerate(cols, start=1):
            value = export_value(row.get(c.key))
            cell = ws.cell(row=excel_row, column=i, value=value if value != "" else None)
            cell.border = border
            if c.kind in XLSX_FORMATS:
                cell.number_format = XLSX_FORMATS[c.kind]
            if r_index == totals_at:
                cell.font = Font(bold=True)
                cell.fill = total_fill
    for i, c in enumerate(cols, start=1):
        longest = max([len(c.label), *(len(str(export_value(r.get(c.key)))) for r in rows[:500])])
        width = 18 if c.kind == "datetime" else min(max(longest + 2, 9), 45)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[head_row].height = 30
    ws.freeze_panes = f"A{head_row + 1}"
    if rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(cols))}{head_row + len(rows)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# Cores do portal (azul da Daticopy e cinzas neutros).
BRAND = colors.HexColor("#156cc4")
BRAND_DARK = colors.HexColor("#143e6a")
BRAND_SOFT = colors.HexColor("#eef6fe")
INK = colors.HexColor("#0f172a")
MUTED = colors.HexColor("#64748b")
LINE = colors.HexColor("#dbe2e8")
ZEBRA = colors.HexColor("#f6f8fa")
TOTAL_BG = colors.HexColor("#e6eef7")
NUMERIC = {"int", "money", "percent", "decimal"}
KEEP_TOGETHER_ROWS = 14  # bloco pequeno não quebra entre páginas
WIDE_COLUMNS = 16  # acima disso (ex.: uma coluna por dia) o PDF sai em A3 deitado, com letra menor


def _logo(height: float) -> Image:
    resource = files("app.assets").joinpath("logo-daticopy.png")
    if not resource.is_file():
        raise RuntimeError("logo do relatório ausente: app/assets/logo-daticopy.png")
    img = Image(io.BytesIO(resource.read_bytes()), height=height, width=height * 4 / 3)
    img.hAlign = "LEFT"
    return img


class _Styles:
    def __init__(self, *, compact: bool = False) -> None:
        size = 6 if compact else 7.5
        self.pad = 2.5 if compact else 6
        self.title = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=INK)
        self.meta = ParagraphStyle("m", fontName="Helvetica", fontSize=8.5, leading=12, textColor=INK)
        self.small = ParagraphStyle("s", fontName="Helvetica", fontSize=7.5, leading=10, textColor=MUTED)
        self.cell = ParagraphStyle("c", fontName="Helvetica", fontSize=size, leading=size + 2, textColor=INK)
        self.cell_r = ParagraphStyle("cr", parent=self.cell, alignment=2)
        self.bold = ParagraphStyle("b", parent=self.cell, fontName="Helvetica-Bold")
        self.bold_r = ParagraphStyle("br", parent=self.bold, alignment=2)
        self.head = ParagraphStyle("h", parent=self.bold, textColor=colors.white)
        self.head_r = ParagraphStyle("hr", parent=self.head, alignment=2)
        self.sec_title = ParagraphStyle(
            "st", fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=BRAND_DARK
        )
        self.sec_detail = ParagraphStyle(
            "sd", fontName="Helvetica", fontSize=7.5, leading=10, textColor=MUTED
        )
        self.stat_label = ParagraphStyle("sl", fontName="Helvetica", fontSize=7.5, leading=9, textColor=MUTED)
        self.stat_value = ParagraphStyle(
            "sv", fontName="Helvetica-Bold", fontSize=14, leading=17, textColor=BRAND_DARK
        )


def _p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


def _header(title: str, meta: list[tuple[str, str]], st: _Styles, width: float) -> Table:
    generated = datetime.now(DISPLAY_TZ).strftime("%d/%m/%Y %H:%M")
    info: list[Any] = [_p(title.upper(), st.title), Spacer(1, 2 * mm)]
    info += [Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", st.meta) for label, value in meta if value]
    right = Paragraph(
        f"Gerado em<br/><b>{escape(generated)}</b>", ParagraphStyle("g", parent=st.small, alignment=2)
    )
    logo_w = 30 * mm
    t = Table(
        [[_logo(22.5 * mm), info, right]], colWidths=[logo_w + 4 * mm, width - logo_w - 44 * mm, 40 * mm]
    )
    t.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("LINEBELOW", (0, 0), (-1, 0), 1.2, BRAND),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 4 * mm),
            ]
        )
    )
    return t


def _summary(stats: list[Stat], st: _Styles, width: float) -> Table:
    cells = [
        [_p(s.label, st.stat_label), _p(text_cell(Col("v", "", s.kind), s.value), st.stat_value)]
        for s in stats
    ]
    t = Table([cells], colWidths=[width / len(stats)] * len(stats))
    style: list[Any] = [
        ("BOX", (0, 0), (-1, -1), 0.6, LINE),
        ("BACKGROUND", (0, 0), (-1, -1), BRAND_SOFT),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]
    style += [("LINEAFTER", (i, 0), (i, 0), 0.6, LINE) for i in range(len(stats) - 1)]
    t.setStyle(TableStyle(style))
    return t


def _grid(
    cols: list[Col],
    rows: list[dict[str, Any]],
    st: _Styles,
    width: float,
    *,
    totals: dict[str, Any] | None = None,
    totals_label: str = "Total",
    lead: list[list[Any]] | None = None,
) -> Table:
    """Table with optional `lead` rows (section header) before the column header."""
    lead = lead or []

    def para(col: Col, value: Any, *, strong: bool = False) -> Paragraph:
        num = col.kind in NUMERIC
        style = (st.bold_r if num else st.bold) if strong else (st.cell_r if num else st.cell)
        return _p(text_cell(col, value), style)

    body: list[list[Any]] = [*lead, [_p(c.label, st.head_r if c.kind in NUMERIC else st.head) for c in cols]]
    head_at = len(lead)
    body += [[para(c, r.get(c.key)) for c in cols] for r in rows]
    if totals is not None:
        body.append(
            [
                _p(totals_label, st.bold)
                if i == 0
                else (para(c, totals.get(c.key), strong=True) if totals.get(c.key) is not None else "")
                for i, c in enumerate(cols)
            ]
        )
    # Texto ocupa mais que número: peso 2,5 para colunas de texto, 1 para as numéricas.
    weights = [1.0 if c.kind in NUMERIC else 2.5 for c in cols]
    unit = width / sum(weights)
    t = Table(body, repeatRows=head_at + 1, colWidths=[w * unit for w in weights])
    style: list[Any] = [
        ("BACKGROUND", (0, head_at), (-1, head_at), BRAND),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, head_at + 1), (-1, -1), 0.4, LINE),
        ("BOX", (0, head_at), (-1, -1), 0.6, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), st.pad),
        ("RIGHTPADDING", (0, 0), (-1, -1), st.pad),
    ]
    for i in range(head_at + 2, len(body), 2):
        style.append(("BACKGROUND", (0, i), (-1, i), ZEBRA))
    if totals is not None:
        style.append(("BACKGROUND", (0, len(body) - 1), (-1, -1), TOTAL_BG))
    for i in range(head_at):
        style += [("SPAN", (0, i), (-1, i)), ("BACKGROUND", (0, i), (-1, i), BRAND_SOFT)]
    if head_at:
        style.append(("BOX", (0, 0), (-1, head_at - 1), 0.6, LINE))
    t.setStyle(TableStyle(style))
    return t


def _sections(data: ReportData, st: _Styles, width: float) -> list[Any]:
    sections = data.sections or {}
    cols = [c for c in data.columns if not c.section]
    out: list[Any] = []
    grouped: dict[str, list[dict[str, Any]]] = {}
    for r in data.rows:
        grouped.setdefault(str(r.get(SECTION_KEY)), []).append(r)
    for key, rows in grouped.items():
        sec = sections.get(key) or Section(key, key)
        details = "   ·   ".join(f"<b>{escape(lb)}:</b> {escape(v)}" for lb, v in sec.details)
        pad: list[Any] = [""] * (len(cols) - 1)
        lead: list[list[Any]] = [
            [_p(sec.title, st.sec_title), *pad],
            [Paragraph(details, st.sec_detail), *pad],
        ]
        table = _grid(cols, rows, st, width, totals=sec.totals, totals_label=sec.totals_label, lead=lead)
        out += [table if len(rows) > KEEP_TOGETHER_ROWS else KeepTogether(table), Spacer(1, 5 * mm)]
    if data.totals:
        out.append(_grid(cols, [], st, width, totals=data.totals, totals_label="Total geral"))
    return out


def _pdf(data: ReportData, *, title: str, meta: list[tuple[str, str]]) -> bytes:
    buf = io.BytesIO()
    visible = [c for c in data.columns if not (data.sections is not None and c.section)]
    wide = len(visible) > WIDE_COLUMNS
    page = landscape(A3 if wide else A4)
    margin = 12 * mm
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=10 * mm,
        bottomMargin=14 * mm,
        title=title,
        author="Daticopy · Dati Monitor",
    )
    st = _Styles(compact=wide)
    width = page[0] - 2 * margin
    story: list[Any] = [_header(title, meta, st, width), Spacer(1, 5 * mm)]
    if data.summary:
        story += [_summary(data.summary, st, width), Spacer(1, 4 * mm)]
    story += [_p(n, st.small) for n in data.notes]
    story.append(Spacer(1, 4 * mm))
    if not data.rows:
        story.append(_p("Nenhum registro com esses filtros.", st.meta))
    elif data.sections is not None:
        story += _sections(data, st, width)
    else:
        story.append(_grid(data.columns, data.rows, st, width, totals=data.totals))

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(LINE)
        canvas.line(margin, 10 * mm, page[0] - margin, 10 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(margin, 6 * mm, f"Daticopy · Dati Monitor · {title}")
        canvas.drawRightString(page[0] - margin, 6 * mm, f"Página {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
